#!/usr/bin/env python3
"""
Dokumenten-Export für TeacherAssist – Wortgutachten als DOCX und PDF.

Der LLM-Skill schreibt Platzhalter ({NAME}, {er/sie}, {sein/ihr}, …) statt
Klarnamen. Erst beim Rendern werden diese hier ersetzt – dies ist die einzige
Stelle im Codepfad, an der der entschlüsselte Klarname in eine Datei fließt
(Datenminimierung, DSGVO Art. 5).

PDF-Pfad:
    1. Primär: docx2pdf (nutzt MS-Word-COM auf Windows → identisches Layout)
    2. Fallback: reportlab (reine Python-PDF mit vereinfachtem Layout +
       Banner-Hinweis, wenn Word nicht verfügbar ist)
"""

import re
from datetime import datetime
from pathlib import Path
from typing import Optional
from xml.sax.saxutils import escape as xml_escape

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.shared import Pt, Cm


# ---------------------------------------------------------------------------
# Pronomen / Platzhalter
# ---------------------------------------------------------------------------

_PRONOUN_MAP = {
    "sie": {"er/sie": "sie", "sein/ihr": "ihr", "ihm/ihr": "ihr", "ihn/sie": "sie"},
    "er":  {"er/sie": "er",  "sein/ihr": "sein", "ihm/ihr": "ihm", "ihn/sie": "ihn"},
}


def _name_reinsertion(text: str, identity: dict) -> str:
    """Ersetzt {NAME}, {er/sie}, {sein/ihr}, {ihm/ihr}, {ihn/sie} (auch capitalized)."""
    vorname = (identity.get("vorname") or identity.get("alias") or "").strip()
    pronomen = (identity.get("pronomen") or "sie").strip().lower()
    table = _PRONOUN_MAP.get(pronomen, _PRONOUN_MAP["sie"])

    out = text.replace("{NAME}", vorname)

    for ph, value in table.items():
        cap = value.capitalize()
        out = out.replace("{" + ph + "}", value)
        if "/" in ph:
            first, second = ph.split("/")
            # akzeptiere {Sein/ihr}, {Sein/Ihr}, {SEIN/IHR}
            out = out.replace("{" + first.capitalize() + "/" + second + "}", cap)
            out = out.replace("{" + first.capitalize() + "/" + second.capitalize() + "}", cap)
            out = out.replace("{" + ph.upper() + "}", cap)
    return out


# ---------------------------------------------------------------------------
# Hilfsformatierung
# ---------------------------------------------------------------------------

def _schuljahr_aus_datum(datum: datetime) -> str:
    """Deutsches Schuljahr (Aug–Jul). Mai 2026 → '2025/26'."""
    if datum.month >= 8:
        start = datum.year
    else:
        start = datum.year - 1
    return f"{start}/{str(start + 1)[-2:]}"


def _format_rating(comp_entry: dict) -> str:
    rating = comp_entry.get("rating")
    scale = comp_entry.get("scale") or ""
    if rating is None:
        return ""
    if scale and "-" in scale:
        max_part = scale.split("-")[-1].strip()
        return f"{rating} / {max_part}"
    return f"{rating}"


def _competency_rows(student: dict) -> list:
    """Liefert [(Fach, Kompetenz, Einstufung, Notiz)]-Tupel, sortiert nach Fach."""
    rows = []
    for fach in sorted(student.get("competencies", {}).keys()):
        for kompetenz in sorted(student["competencies"][fach].keys()):
            entry = student["competencies"][fach][kompetenz]
            rows.append((fach, kompetenz, _format_rating(entry), entry.get("note", "")))
    return rows


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def generate_wortgutachten_docx(
    student: dict,
    narrative: str,
    out_path: Path,
    schuljahr: Optional[str] = None,
    datum: Optional[str] = None,
) -> Path:
    """Erzeugt ein Wortgutachten als .docx. narrative kann Platzhalter enthalten."""
    identity = student.get("identity", {})
    vorname = identity.get("vorname", "")
    nachname = identity.get("nachname", "")
    klasse = student.get("klasse", "")

    dt = datetime.fromisoformat(datum) if datum else datetime.now()
    if schuljahr is None:
        schuljahr = _schuljahr_aus_datum(dt)
    datum_str = dt.strftime("%d.%m.%Y")

    narrative_rendered = _name_reinsertion(narrative, identity)

    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_paragraph()
    run = title.add_run(f"Wortgutachten – {vorname} {nachname}".strip(" –"))
    run.bold = True
    run.font.size = Pt(16)

    meta = doc.add_paragraph()
    meta.add_run(f"Klasse {klasse}   ·   Schuljahr {schuljahr}   ·   {datum_str}").italic = True

    rows = _competency_rows(student)
    if rows:
        doc.add_paragraph().add_run("Kompetenzen").bold = True
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, label in enumerate(("Fach", "Kompetenz", "Einstufung", "Notiz")):
            r = hdr[i].paragraphs[0].add_run(label)
            r.bold = True
            hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for fach, kompetenz, einstufung, notiz in rows:
            row = table.add_row().cells
            row[0].text = fach
            row[1].text = kompetenz
            row[2].text = einstufung
            row[3].text = notiz

    doc.add_paragraph()
    doc.add_paragraph().add_run("Verbale Einschätzung").bold = True
    for absatz in re.split(r"\n{2,}", narrative_rendered.strip()):
        absatz = absatz.strip()
        if not absatz:
            continue
        para = doc.add_paragraph(absatz)
        para.paragraph_format.space_after = Pt(8)

    doc.add_paragraph()
    sig = doc.add_paragraph()
    sig.add_run(f"{datum_str}                                                                                                ").font.size = Pt(10)
    sig.add_run("Unterschrift Lehrkraft").italic = True

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def generate_wortgutachten_pdf(
    student: dict,
    narrative: str,
    out_path: Path,
    schuljahr: Optional[str] = None,
    datum: Optional[str] = None,
) -> Path:
    """Erzeugt ein Wortgutachten als .pdf. Primär via docx2pdf (Word), Fallback reportlab."""
    out_path.parent.mkdir(parents=True, exist_ok=True)

    try:
        return _pdf_via_docx2pdf(student, narrative, out_path, schuljahr, datum)
    except Exception:
        return _pdf_via_reportlab(student, narrative, out_path, schuljahr, datum)


def _pdf_via_docx2pdf(student, narrative, out_path, schuljahr, datum) -> Path:
    """DOCX → PDF via MS-Word-COM. Wirft, wenn Word nicht installiert ist."""
    from docx2pdf import convert
    tmp_docx = out_path.with_suffix(".tmp.docx")
    try:
        generate_wortgutachten_docx(student, narrative, tmp_docx, schuljahr, datum)
        convert(str(tmp_docx), str(out_path))
        return out_path
    finally:
        if tmp_docx.exists():
            tmp_docx.unlink()


def _pdf_via_reportlab(student, narrative, out_path, schuljahr, datum) -> Path:
    """Reine Python-PDF mit vereinfachtem Layout + Hinweisbanner."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    )
    from reportlab.lib import colors

    identity = student.get("identity", {})
    vorname = identity.get("vorname", "")
    nachname = identity.get("nachname", "")
    klasse = student.get("klasse", "")

    dt = datetime.fromisoformat(datum) if datum else datetime.now()
    if schuljahr is None:
        schuljahr = _schuljahr_aus_datum(dt)
    datum_str = dt.strftime("%d.%m.%Y")

    narrative_rendered = _name_reinsertion(narrative, identity)

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], fontSize=16, alignment=0, spaceAfter=4)
    meta = ParagraphStyle("meta", parent=styles["Normal"], fontSize=10, textColor=colors.grey, spaceAfter=18)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=12, spaceBefore=10, spaceAfter=6)
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=11, spaceAfter=8, leading=15)
    banner = ParagraphStyle("banner", parent=styles["Italic"], fontSize=9, textColor=colors.HexColor("#a85a00"))

    doc = SimpleDocTemplate(
        str(out_path), pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=1.6 * cm, bottomMargin=1.6 * cm,
    )

    story = [
        Paragraph("Hinweis: Erstellt ohne MS Word – Layout vereinfacht.", banner),
        Spacer(1, 4),
        Paragraph(xml_escape(f"Wortgutachten – {vorname} {nachname}".strip(" –")), h1),
        Paragraph(f"Klasse {xml_escape(klasse)} &nbsp;·&nbsp; Schuljahr {xml_escape(schuljahr)} &nbsp;·&nbsp; {xml_escape(datum_str)}", meta),
    ]

    rows = _competency_rows(student)
    if rows:
        story.append(Paragraph("Kompetenzen", h2))
        data = [["Fach", "Kompetenz", "Einstufung", "Notiz"]] + [list(r) for r in rows]
        table = Table(data, colWidths=[3.5 * cm, 5.5 * cm, 2.5 * cm, 5.5 * cm])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e7eef7")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(table)
        story.append(Spacer(1, 12))

    story.append(Paragraph("Verbale Einschätzung", h2))
    for absatz in re.split(r"\n{2,}", narrative_rendered.strip()):
        absatz = absatz.strip()
        if not absatz:
            continue
        story.append(Paragraph(xml_escape(absatz).replace("\n", "<br/>"), body))

    story.append(Spacer(1, 24))
    story.append(Paragraph(
        f"{xml_escape(datum_str)} &nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; "
        "<i>Unterschrift Lehrkraft</i>",
        body,
    ))

    doc.build(story)
    return out_path
